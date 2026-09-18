#!/usr/bin/env python3
"""Build a complete IJCM submission set, parallel to Study 1 / Data_IJCM."""
from __future__ import annotations

import csv
import hashlib
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
STUDY = ROOT / "Study4_FINAL_VERSION"
SOURCE = STUDY / "manuscript" / "Study4_Manuscript.md"
OUT = ROOT / "Study4_IJCM_Submission"
DATA_ROOT = ROOT / "Data_Study4_IJCM"
TITLE = (
    "How Does the Civil Engineering Industry Economy Change under an AI Shock? "
    "The United Kingdom’s Service Links with India and China"
)
RUNNING_HEAD = "AI and the civil engineering economy"
AUTHOR = "Yuxuan Chai"
AFFILIATION = "University of Strathclyde, Glasgow, United Kingdom"
EMAIL = "yuxuanchai98@outlook.com"
JOURNAL = "International Journal of Construction Management"

MAIN_FIGURES = [
    ("Figure1.png", "figure22_uk_india_china_sj3.png",
     "The United Kingdom’s SJ3 service corridors with India and China, 2015–2024 (2019 = 100)."),
    ("Figure2.png", "figure23_uk_india_china_service_growth.png",
     "Growth in UK–India and UK–China services trade by category, 2019–2024."),
    ("Figure3.png", "figure12_tnlg_vs_tml.png",
     "EU-27 professional-service NLG adoption and technology comparisons."),
    ("Figure4.png", "figure15_cross_section_tnlg.png",
     "Importer ΔTNLG and the change in Indian SJ3 imports."),
    ("Figure5.png", "figure17_sbs_turnover_m71_vs_F.png",
     "M71 and construction turnover growth, 2021–2024."),
]

ABSTRACT = """This paper studies the civil-engineering economy under a generative-AI shock through two United Kingdom-centred relationships: UK–India and UK–China. Bilateral evidence covers technical and other business services (SJ3), construction services (SE), and computer and information services (SI). A panel of 17 EU importers supplies harmonised variation in enterprise use of natural-language-generation (NLG) AI. UK trends are descriptive; the panel tests whether importer adoption is associated with Indian Mode-1 services or Chinese project-based construction services. The shock is the 2023–24 change in Eurostat NLG use among professional-service enterprises. Domestic outcomes are turnover, employment, wages and value added in NACE M71 architectural and engineering activities; trade outcomes come from OECD–WTO BaTIS. Between 2019 and 2024, UK imports of Indian SJ3 rose 127.7% and UK imports of Chinese SJ3 rose 99.3%. In the importer panel, the India-only SJ3 coefficient is 0.021 (s.e. 0.008), whereas Chinese construction services are unrelated to importer NLG adoption (−0.008, s.e. 0.009). Domestic M71 turnover and wages grew more slowly in higher-adoption economies, while employment was unchanged. The evidence is consistent with generative tools changing the coordination of digitally deliverable services, especially between the UK and India. It does not show that AI caused the bilateral changes or that civil-engineering jobs or GDP relocated."""

KEYWORDS = (
    "generative AI; civil engineering; construction management; services trade; "
    "NACE M71; Mode 1"
)

TABLES = [
    {
        "title": "UK–India and UK–China service corridors",
        "headers": ["Import flow", "Service", "2019 (USD m)", "2024 (USD m)", "Change"],
        "rows": [
            ["UK imports from India", "SJ3 technical and other business services", "2,186.6", "4,979.3", "127.7%"],
            ["India imports from UK", "SJ3 technical and other business services", "787.6", "1,554.2", "97.3%"],
            ["UK imports from China", "SJ3 technical and other business services", "775.5", "1,546.0", "99.3%"],
            ["China imports from UK", "SJ3 technical and other business services", "739.0", "954.6", "29.2%"],
            ["UK imports from India", "SE construction services", "41.9", "72.8", "73.7%"],
            ["India imports from UK", "SE construction services", "39.3", "72.6", "84.8%"],
            ["UK imports from China", "SE construction services", "118.5", "179.8", "51.7%"],
            ["China imports from UK", "SE construction services", "54.8", "70.8", "29.2%"],
            ["UK imports from India", "SI computer and information services", "2,215.3", "4,723.0", "113.2%"],
            ["India imports from UK", "SI computer and information services", "389.2", "707.6", "81.8%"],
            ["UK imports from China", "SI computer and information services", "724.7", "1,054.5", "45.5%"],
            ["China imports from UK", "SI computer and information services", "1,639.3", "2,404.1", "46.7%"],
        ],
        "note": "Source: OECD–WTO BaTIS, adjustment B. Values are balanced imports in current USD million.",
    },
    {
        "title": "EU-27 enterprise use of natural-language-generation AI",
        "headers": ["Industry", "2023", "2024", "2025"],
        "rows": [
            ["NACE M professional services", "4.55%", "11.51%", "17.74%"],
            ["NACE F construction", "0.58%", "2.42%", "3.25%"],
        ],
        "note": "Source: Eurostat isoc_eb_ain2, enterprises with at least ten persons.",
    },
    {
        "title": "Selected trade estimates under the NLG shock",
        "headers": ["Outcome/specification", "Coefficient (s.e.)", "N"],
        "rows": [
            ["Indian SJ3, Post × ΔM TNLG", "0.021*** (0.008)", "119"],
            ["Pooled Mode-1 SJ3, Post-2024 × ΔM TNLG", "0.007** (0.003)", "357"],
            ["Chinese construction services (SE)", "−0.008 (0.009)", "119"],
            ["Chinese SJ3", "0.005 (0.010)", "119"],
            ["Indian computer services comparison", "−0.002 (0.006)", "357"],
            ["Generic any-AI comparison", "0.000 (0.006)", "336"],
        ],
        "note": "Importer and year fixed effects; standard errors clustered by importer. The pooled model also contains partner fixed effects. *p < 0.10, **p < 0.05, ***p < 0.01.",
    },
    {
        "title": "M71 industry outcomes under the NLG shock",
        "headers": ["Outcome", "Post-2024 × ΔTNLG (s.e.)", "N"],
        "rows": [
            ["Log M71 turnover", "−0.007** (0.003)", "68"],
            ["Log M71 employment", "−0.000 (0.002)", "68"],
            ["Log M71 wages", "−0.009*** (0.003)", "68"],
            ["Log M71 value added", "−0.006*** (0.002)", "68"],
            ["Log construction turnover (comparison)", "−0.010*** (0.003)", "68"],
        ],
        "note": "Country and year fixed effects; standard errors clustered by country. *p < 0.10, **p < 0.05, ***p < 0.01.",
    },
    {
        "title": "Chinese construction services by importer NLG group",
        "headers": ["EU destination group", "2019 (USD m)", "2024 (USD m)", "Change"],
        "rows": [
            ["Lower ΔTNLG", "571.9", "946.2", "65.4%"],
            ["Higher ΔTNLG", "679.9", "1,002.6", "47.5%"],
        ],
        "note": "Source: OECD–WTO BaTIS SE. Groups split at the importer median change in NACE M TNLG.",
    },
    {
        "title": "Interpretation boundary",
        "headers": ["Supported by the data", "Not identified"],
        "rows": [
            ["Two-way UK–India and UK–China service trends", "AI causing those bilateral changes"],
            ["EU importer NLG associated with Indian SJ3", "Civil-specific SJ312 invoices"],
            ["Null Chinese construction-service comparison", "India replacing Chinese or UK civil engineers"],
            ["M71 outcomes associated with importer NLG", "Bilateral employment or GDP effects"],
        ],
        "note": "The table restates the claim boundary; it is not an additional statistical test.",
    },
]

REFERENCES = """
Acemoglu, D., & Restrepo, P. (2018). The race between man and machine: Implications of technology for growth, factor shares, and employment. *American Economic Review, 108*(6), 1488–1542.

Acemoglu, D., & Restrepo, P. (2019). Automation and new tasks: How technology displaces and reinstates labor. *Journal of Economic Perspectives, 33*(2), 3–30.

Acemoglu, D., & Restrepo, P. (2020). Robots and jobs: Evidence from US labor markets. *Journal of Political Economy, 128*(6), 2188–2244.

Agrawal, A., Gans, J., & Goldfarb, A. (2018). *Prediction machines: The simple economics of artificial intelligence*. Harvard Business Review Press.

Agrawal, A., Gans, J., & Goldfarb, A. (2022). *Power and prediction: The disruptive economics of artificial intelligence*. Harvard Business Review Press.

Autor, D. H. (2015). Why are there still so many jobs? The history and future of workplace automation. *Journal of Economic Perspectives, 29*(3), 3–30.

Autor, D. H., Levy, F., & Murnane, R. J. (2003). The skill content of recent technological change: An empirical exploration. *Quarterly Journal of Economics, 118*(4), 1279–1333.

Baldwin, R. (2016). *The great convergence: Information technology and the new globalization*. Harvard University Press.

Baldwin, R. (2019). *The globotics upheaval: Globalization, robotics, and the future of work*. Oxford University Press.

Blinder, A. S. (2006). Offshoring: The next industrial revolution? *Foreign Affairs, 85*(2), 113–128.

Brynjolfsson, E., Li, D., & Raymond, L. (2025). Generative AI at work. *Quarterly Journal of Economics, 140*(2), 889–942.

Dell’Acqua, F., McFowland, E., Mollick, E., Lifshitz-Assaf, H., Kellogg, K. C., Rajendran, S., Krayer, L., Candelon, F., & Lakhani, K. R. (2023). Navigating the jagged technological frontier: Field experimental evidence of the effects of AI on knowledge worker productivity and quality (Harvard Business School Working Paper 24-013).

Eastman, C., Teicholz, P., Sacks, R., & Liston, K. (2011). *BIM handbook: A guide to building information modeling* (2nd ed.). Wiley.

Eloundou, T., Manning, S., Mishkin, P., & Rock, D. (2023). GPTs are GPTs: An early look at the labor market impact potential of large language models. arXiv:2303.10130.

Eloundou, T., Manning, S., Mishkin, P., & Rock, D. (2024). GPTs are GPTs: Labor market impact potential of LLMs. *Science, 384*(6702), 1306–1308.

Eurostat. (2024). *ICT usage in enterprises: Artificial intelligence (isoc_eb_ain2); national accounts (nama_10_a64); structural business statistics (sbs_ovw_act, sbs_sc_ovw)*. European Commission.

Felten, E., Raj, M., & Seamans, R. (2018). A method to link advances in artificial intelligence to occupational abilities. *AEA Papers and Proceedings, 108*, 54–57.

Felten, E., Raj, M., & Seamans, R. (2021). Occupational, industry, and geographic exposure to artificial intelligence: A novel dataset and its potential uses. *Strategic Management Journal, 42*(12), 2195–2217.

Fortanier, F., Liberatore, A., Maurer, A., & Pilgrim, G. (2017). *The OECD–WTO Balanced Trade in Services database*. OECD/WTO.

Francois, J., & Hoekman, B. (2010). Services trade and policy. *Journal of Economic Literature, 48*(3), 642–692.

Gann, D. M., & Salter, A. J. (2000). Innovation in project-based, service-enhanced firms: The construction of complex products and systems. *Research Policy, 29*(7–8), 955–972.

Grossman, G. M., & Rossi-Hansberg, E. (2008). Trading tasks: A simple theory of offshoring. *American Economic Review, 98*(5), 1978–1997.

Hui, X., Reshef, O., & Zhou, L. (2024). The short-term effects of generative artificial intelligence on employment: Evidence from an online labor market. *Organization Science*.

ILO. (n.d.). *ILOSTAT: Employment by sex and economic activity*. International Labour Organization.

Korinek, A., & Stiglitz, J. E. (2021). Artificial intelligence, globalization, and strategies for economic development (NBER Working Paper 28453).

Loungani, P., Mishra, S., Papageorgiou, C., & Wang, K. (2017). World trade in services: Evidence from a new dataset (IMF Working Paper 17/77).

Noy, S., & Zhang, W. (2023). Experimental evidence on the productivity effects of generative artificial intelligence. *Science, 381*(6654), 187–192.

OECD & WTO. (n.d.). *BaTIS: Balanced Trade in Services dataset* (BPM6, adjustment B).

Office for National Statistics. (2019). *Which occupations are at highest risk of being automated?* (Table 9, SOC 2010).

Peng, S., Kalliamvakou, E., Cihon, P., & Demirer, M. (2023). The impact of AI on developer productivity: Evidence from GitHub Copilot. arXiv:2302.06590.

Sacks, R., Eastman, C., Lee, G., & Teicholz, P. (2018). *BIM handbook: A guide to building information modeling for owners, designers, engineers, contractors, and facility managers* (3rd ed.). Wiley.

Webb, M. (2020). *The impact of artificial intelligence on the labor market*. Stanford University.

Whyte, J. (2019). How digital information transforms project delivery models. *Project Management Journal, 50*(2), 177–194.

Winch, G. M. (2010). *Managing construction projects* (2nd ed.). Wiley-Blackwell.

WTO. (1994). *General Agreement on Trade in Services*. World Trade Organization.
""".strip()


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_document(document: Document, *, double: bool = True) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.ONE_POINT_FIVE
    pf.line_spacing = 2.0 if double else 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    for name, size, bold in [
        ("Title", 16, True),
        ("Heading 1", 12, True),
        ("Heading 2", 12, True),
        ("Heading 3", 12, True),
    ]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2.0 if double else 1.5
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, size=10)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def add_line_numbering(section) -> None:
    sect_pr = section._sectPr
    for child in list(sect_pr):
        if child.tag == qn("w:lnNumType"):
            sect_pr.remove(child)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "newPage")
    sect_pr.append(ln)


def add_header(section, text: str) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    set_run_font(run, size=10, italic=True)


def make_reference_doc(path: Path) -> None:
    document = Document()
    style_document(document, double=True)
    document.save(path)


def pandoc(source: Path, dest: Path, reference: Path, resource: Path) -> None:
    subprocess.check_call(
        [
            "pandoc",
            str(source),
            "--from=markdown+tex_math_dollars",
            "--to=docx",
            "--output",
            str(dest),
            f"--resource-path={resource}:{STUDY}:{STUDY / 'manuscript'}:{STUDY / 'figures'}",
            f"--reference-doc={reference}",
        ]
    )


def polish_docx(path: Path, *, line_numbers: bool = True, header: str | None = None) -> None:
    document = Document(str(path))
    style_document(document, double=True)
    section = document.sections[0]
    add_page_number(section)
    if line_numbers:
        add_line_numbering(section)
    if header:
        add_header(section, header)
    if document.paragraphs:
        title_p = document.paragraphs[0]
        title_p.style = document.styles["Title"]
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.first_line_indent = Cm(0)
        title_p.paragraph_format.space_after = Pt(12)
        for run in title_p.runs:
            set_run_font(run, size=16, bold=True)
    for para in document.paragraphs[1:]:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if style_name.startswith("Heading") or style_name == "Title":
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.line_spacing = 2.0
            for run in para.runs:
                set_run_font(run, size=12 if style_name != "Title" else 16, bold=True)
            continue
        para.paragraph_format.line_spacing = 2.0
        para.paragraph_format.space_after = Pt(0)
        if not text:
            para.paragraph_format.first_line_indent = Cm(0)
            continue
        if text.lower() in {"abstract", "keywords", "references", "tables", "figure captions",
                            "disclosure statement", "data availability statement"}:
            para.paragraph_format.first_line_indent = Cm(0)
            for run in para.runs:
                set_run_font(run, size=12, bold=True)
            continue
        if text.startswith(("Table ", "Figure ", "Note.", "Source:", "Keywords:")):
            para.paragraph_format.first_line_indent = Cm(0)
            continue
        if style_name in {"Normal", "Body Text", "First Paragraph"}:
            para.paragraph_format.first_line_indent = Cm(1.27)
        for run in para.runs:
            if run.font.size is None:
                set_run_font(run, size=12)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing = 1.15
                    para.paragraph_format.first_line_indent = Cm(0)
                    for run in para.runs:
                        set_run_font(run, size=10)
    for shape in document.inline_shapes:
        try:
            width = shape.width
            height = shape.height
            max_w = Cm(15.5)
            if width > max_w:
                ratio = max_w / width
                shape.width = int(width * ratio)
                shape.height = int(height * ratio)
        except Exception:
            pass
    document.save(path)


def markdown_tables_and_figures() -> str:
    blocks = ["## Tables\n"]
    for i, table in enumerate(TABLES, 1):
        blocks.append(f"**Table {i}.** {table['title']}\n")
        header = "| " + " | ".join(table["headers"]) + " |"
        sep = "|" + "|".join(["---"] * len(table["headers"])) + "|"
        rows = ["| " + " | ".join(row) + " |" for row in table["rows"]]
        blocks.append("\n".join([header, sep, *rows]))
        blocks.append(f"\n*Note.* {table['note']}\n")
    blocks.append("## Figure captions\n")
    for i, (_, _, caption) in enumerate(MAIN_FIGURES, 1):
        blocks.append(f"**Figure {i}.** {caption}\n")
    blocks.append("## Figures\n")
    blocks.append(
        "Figures are supplied as separate 300 dpi PNG files (`Figure1.png`–`Figure5.png`) "
        "and are repeated below for the convenience of reviewers.\n"
    )
    for i, (_, source_name, caption) in enumerate(MAIN_FIGURES, 1):
        blocks.append(f"**Figure {i}.** {caption}\n")
        blocks.append(f"![Figure {i}](figures/{source_name})\n")
    return "\n".join(blocks)


def format_source_manuscript() -> str:
    text = SOURCE.read_text(encoding="utf-8")
    start = text.index("## 1. Introduction")
    end = text.index("## References")
    body = text[start:end].rstrip()
    body = body.replace(
        "Replication files are in `Study4_Article_Complete/` "
        "(open `Study4_Manuscript.md`, `CN_full_article.md`, `FIGURES.md`, "
        "and `tables_for_article.md` in the editor — not the zip). Figures, "
        "tables, and official data are in `figures/`, `tables/`, and `data/` "
        "in that folder. Missing series are catalogued in `DATA_INVENTORY.md`; "
        "their absence is part of the claim boundary.",
        "A complete replication package accompanies this submission. Missing "
        "series are catalogued there; their absence is part of the claim boundary.",
    )
    body = body.replace(
        "All files are in `data/` of the replication package. Retrieval dates are in `DATA_SOURCES.md`. Table A in the inventory file records failed pulls (OWID ChatGPT CSVs, IMF AIPI empty JSON, OECD ICT_BUS 404, ILO M71 404). Failed pulls are not filled.",
        "The analysis uses retrieved official series only. No country, industry or trade cell is interpolated. Retrieval dates and unsuccessful pulls (including OWID ChatGPT files, an empty IMF AI Preparedness response, OECD ICT_BUS, and ILO M71) are documented in the supplementary replication package and are not filled.",
    )
    body = body.replace(
        "Korinek and Stiglitz (2021) and Acemoglu (2025, policy writing on complementary AI) warn that the distribution of gains depends on whether AI automates or augments.",
        "Korinek and Stiglitz (2021) warn that the distribution of gains depends on whether AI automates or augments.",
    )
    body = body.replace(
        "Whyte and colleagues on digital delivery similarly locate the information model in professional organisations.",
        "Whyte (2019) similarly locates digital delivery in professional organisations rather than only on the construction site.",
    )
    header = (
        f"# {TITLE}\n\n"
        f"{AUTHOR}\n\n"
        f"{AFFILIATION}\n\n"
        f"Correspondence: {AUTHOR}, {AFFILIATION}. Email: {EMAIL}\n\n"
        f"Manuscript submitted to the *{JOURNAL}*.\n\n"
        "## Abstract\n\n"
        f"{ABSTRACT}\n\n"
        f"**Keywords:** {KEYWORDS}\n\n"
    )
    back = (
        "\n\n## Disclosure statement\n\n"
        "No potential conflict of interest was reported by the author.\n\n"
        "## Data availability statement\n\n"
        "The data that support the findings of this study are included in the "
        "accompanying replication package (`Data_Study4_IJCM/`). Eurostat, "
        "OECD–WTO BaTIS, ILOSTAT and World Bank series remain subject to their "
        "publishers’ terms and should be cited as in the reference list. "
        "Author-generated files are released under CC BY 4.0.\n\n"
        "## References\n\n"
        f"{REFERENCES}\n\n"
        f"{markdown_tables_and_figures()}"
    )
    identified = header + body + back
    SOURCE.write_text(identified, encoding="utf-8")
    (STUDY / "manuscript" / "article_body.md").write_text(identified, encoding="utf-8")
    return identified


def blinded_markdown(identified: str) -> str:
    text = identified
    text = re.sub(rf"^{re.escape(AUTHOR)}\n\n", "", text, count=1, flags=re.M)
    text = text.replace(f"{AFFILIATION}\n\n", "")
    text = text.replace(
        f"Correspondence: {AUTHOR}, {AFFILIATION}. Email: {EMAIL}\n\n",
        "",
    )
    text = text.replace(f"Manuscript submitted to the *{JOURNAL}*.\n\n", "")
    text = text.replace("(`Data_Study4_IJCM/`)", "(supplied as supplementary material)")
    text = re.sub(r"(?im)^.*(?:Yuxuan Chai|Strathclyde|yuxuanchai98@outlook\.com).*\n?", "", text)
    return text


def add_styled_paragraph(document: Document, text: str, *, bold=False, italic=False,
                         size=12, align="left", space_after=0, indent=False):
    p = document.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(1.27) if indent else Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def write_title_page(path: Path, word_count: int) -> None:
    document = Document()
    style_document(document, double=True)
    add_page_number(document.sections[0])
    add_header(document.sections[0], RUNNING_HEAD)
    add_styled_paragraph(document, JOURNAL, italic=True, size=12, align="center")
    add_styled_paragraph(document, "Original research article", size=12, align="center")
    document.add_paragraph()
    add_styled_paragraph(document, TITLE, bold=True, size=16, align="center")
    document.add_paragraph()
    add_styled_paragraph(document, AUTHOR, bold=True, size=12, align="center")
    add_styled_paragraph(document, AFFILIATION, size=12, align="center")
    add_styled_paragraph(
        document,
        f"Corresponding author: {AUTHOR}, {AFFILIATION}. Email: {EMAIL}",
        size=12,
        align="center",
    )
    document.add_paragraph()
    add_styled_paragraph(document, "Manuscript details", bold=True, size=12)
    details = [
        f"Running head: {RUNNING_HEAD}",
        f"Word count (including abstract, tables, captions and references): {word_count}",
        "Tables: 6",
        "Figures: 5 (separate 300 dpi PNG files)",
        "Supplementary material: numbered replication package Data_Study4_IJCM",
    ]
    for line in details:
        add_styled_paragraph(document, line, size=12)
    document.add_paragraph()
    add_styled_paragraph(document, "Declarations", bold=True, size=12)
    add_styled_paragraph(
        document,
        "Author contribution: Yuxuan Chai is the sole author of this manuscript.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "Disclosure statement: No potential conflict of interest was reported by the author.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "Funding: This research received no specific grant from any funding agency in the "
        "public, commercial, or not-for-profit sectors.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "Ethics: The study analyses published aggregate statistical series and does not "
        "involve human participants, personal data or animal subjects.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "Data availability: The replication files that support the findings are in "
        "Data_Study4_IJCM. Third-party Eurostat, OECD–WTO BaTIS, ILOSTAT and World Bank "
        "material remains subject to the publishers’ terms.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "This title page is not for peer review. The anonymous manuscript is supplied separately.",
        italic=True,
        size=12,
    )
    document.save(path)


def write_cover_letter(path: Path) -> None:
    document = Document()
    style_document(document, double=True)
    add_page_number(document.sections[0])
    add_styled_paragraph(document, AUTHOR, size=12)
    add_styled_paragraph(document, AFFILIATION, size=12)
    add_styled_paragraph(document, EMAIL, size=12)
    add_styled_paragraph(document, "18 September 2026", size=12)
    document.add_paragraph()
    add_styled_paragraph(document, f"Editor-in-Chief, {JOURNAL}", size=12)
    document.add_paragraph()
    add_styled_paragraph(document, "Dear Editor,", size=12)
    add_styled_paragraph(
        document,
        f"Please consider the enclosed original research article, “{TITLE}”, for publication "
        f"in the {JOURNAL}.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "The manuscript examines how the initial diffusion of natural-language-generation AI "
        "is associated with architectural and engineering activity (NACE M71) and with the "
        "United Kingdom’s services relationships with India and China. It distinguishes "
        "engineering consultancies from construction contractors, and technical-business, "
        "construction and computer service flows. The evidence uses retrieved Eurostat and "
        "OECD–WTO BaTIS series. No missing country, industry or trade cell is interpolated.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(
        document,
        "I confirm that the work is original, has not been published, and is not under "
        "consideration elsewhere. I am the sole author and approve this submission. There "
        "are no competing interests. The research received no specific grant from any "
        "funding agency. A complete replication package is supplied as supplementary material.",
        size=12,
        indent=True,
    )
    add_styled_paragraph(document, "Yours sincerely,", size=12)
    document.add_paragraph()
    add_styled_paragraph(document, AUTHOR, size=12)
    add_styled_paragraph(document, AFFILIATION, size=12)
    add_styled_paragraph(document, EMAIL, size=12)
    document.save(path)


def write_tables_docx(path: Path) -> None:
    document = Document()
    style_document(document, double=True)
    add_page_number(document.sections[0])
    add_header(document.sections[0], "Tables")
    add_styled_paragraph(document, "Tables", bold=True, size=14, align="center")
    add_styled_paragraph(
        document,
        "Editable tables for the International Journal of Construction Management. "
        "Place each table after the paragraph in which it is first cited, or upload "
        "this file as a separate tables document if required by ScholarOne.",
        italic=True,
        size=12,
    )
    for i, table in enumerate(TABLES, 1):
        add_styled_paragraph(document, f"Table {i}. {table['title']}", bold=True, size=12)
        tbl = document.add_table(rows=1 + len(table["rows"]), cols=len(table["headers"]))
        tbl.style = "Table Grid"
        for j, header in enumerate(table["headers"]):
            cell = tbl.rows[0].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(header)
            set_run_font(run, size=10, bold=True)
        for r, row in enumerate(table["rows"], 1):
            for c, value in enumerate(row):
                cell = tbl.rows[r].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(value)
                set_run_font(run, size=10)
        add_styled_paragraph(document, f"Note. {table['note']}", italic=True, size=10)
        document.add_paragraph()
    document.save(path)


def write_figure_captions_docx(path: Path) -> None:
    document = Document()
    style_document(document, double=True)
    add_page_number(document.sections[0])
    add_styled_paragraph(document, "Figure captions", bold=True, size=14, align="center")
    for i, (filename, _, caption) in enumerate(MAIN_FIGURES, 1):
        add_styled_paragraph(document, f"Figure {i}. {caption}", size=12)
        add_styled_paragraph(document, f"File: {filename} (300 dpi PNG, anonymised).", italic=True, size=11)
    document.save(path)


def write_table_csvs(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for i, table in enumerate(TABLES, 1):
        path = folder / f"Table{i}.csv"
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(table["headers"])
            writer.writerows(table["rows"])
            writer.writerow([])
            writer.writerow(["Note", table["note"]])


def export_figures(dest: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    for filename, source_name, _ in MAIN_FIGURES:
        image = Image.open(STUDY / "figures" / source_name)
        dpi = image.info.get("dpi", (200, 200))[0] or 200
        scale = max(1.0, 300 / float(dpi))
        if scale > 1:
            image = image.resize(
                (round(image.width * scale), round(image.height * scale)),
                Image.Resampling.LANCZOS,
            )
        rgb = image.convert("RGB")
        rgb.save(dest / filename, dpi=(300, 300), optimize=True)


def copy_file(src: Path, dst: Path) -> None:
    if src.is_file():
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)


def build_data_package(dest: Path) -> None:
    if dest.exists():
        shutil.rmtree(dest)
    mapping = {
        "01_ai_shock": [
            "eurostat_ai_raw.csv",
            "eurostat_ai_genai_types.csv",
            "eurostat_ai_nace_placebos.csv",
            "owid_fetch_log.csv",
            "imf_aipi_fetch_log.csv",
            "imf_aipi_raw.json",
            "oecd_ict_fetch_log.csv",
            "ms_ai_diffusion_q1_2026.csv",
        ],
        "02_industry_accounts": [
            "eurostat_sbs_M71_F_M.csv",
            "eurostat_nama_gva_M71_F_M.csv",
            "eurostat_nama_emp_M71_F_M.csv",
            "eurostat_nama_D1_P1_M71_F_M.csv",
            "ilo_emp_FM.csv",
            "ilo_m71_fetch_log.csv",
            "bls_engineering_ces.csv",
            "bls_54133_fetch_log.csv",
        ],
        "03_trade": [
            "batis_uk_india_china.csv",
            "batis_civil_related.csv",
            "batis_SJ1_SJ2.csv",
            "panel_eu_sj3_tnlg.csv",
            "panel_eu_sj3_mode1.csv",
        ],
        "04_external_indices": [
            "eloundou_occ_level.csv",
        ],
    }
    data = STUDY / "data"
    for folder, names in mapping.items():
        for name in names:
            copy_file(data / name, dest / folder / name)
    for src_dir, label in [
        (data / "from_legacy_study4", "from_legacy_study4"),
        (data / "from_study1", "from_study1"),
    ]:
        if src_dir.is_dir():
            for path in src_dir.iterdir():
                if path.is_file():
                    copy_file(path, dest / "04_external_indices" / label / path.name)
    derived_files = [
        STUDY / "results.json",
        STUDY / "results_nlg.json",
        STUDY / "results_industry.json",
        STUDY / "tables" / "table_nlg_identification.csv",
        STUDY / "tables" / "table_industry_economy.csv",
        STUDY / "tables" / "table_uk_india_china_corridors.csv",
        STUDY / "tables" / "table_china_se_by_tnlg_split.csv",
        STUDY / "tables" / "table_tnlg_eu27.csv",
        STUDY / "tables" / "table_event_study_tnlg.csv",
        STUDY / "tables" / "table_cross_section_india_tnlg.csv",
        STUDY / "tables" / "table_sbs_m71_growth.csv",
    ]
    for path in derived_files:
        copy_file(path, dest / "05_derived" / path.name)
    write_table_csvs(dest / "06_tables")
    script_names = [
        "download_ai_shocks.py",
        "download_m71_economy.py",
        "run_focused_corridors.py",
        "run_nlg_shock.py",
        "run_industry_economy.py",
        "run_analysis.py",
        "run_novelty_layer.py",
        "run_all_experiments.py",
        "inherit_legacy_instruments.py",
        "build_article_assets.py",
        "build_focused_main_text.py",
        "build_ijcm_submission.py",
    ]
    for name in script_names:
        copy_file(STUDY / "scripts" / name, dest / "07_scripts" / name)
    export_figures(dest / "08_figures")
    copy_file(STUDY / "DATA_SOURCES.md", dest / "DATA_SOURCES.md")
    copy_file(STUDY / "DATA_INVENTORY.md", dest / "DATA_INVENTORY.md")
    (dest / "README.txt").write_text(
        f"""Data accompanying
{TITLE}

Manuscript submitted to {JOURNAL}.
Corresponding author: {AUTHOR} ({EMAIL}), University of Strathclyde.

The article uses retrieved official series. No country, industry or trade cell is interpolated.
Preferred shock: Eurostat E_AI_TNLG, NACE M, 2024 minus 2023.
Industry outcomes: Eurostat NACE M71. Trade: OECD–WTO BaTIS, adjustment B.
Country focus: UK–India and UK–China; EU-17 is the estimation panel only.

Folder contents
---------------
01_ai_shock/
  Eurostat enterprise AI series (TANY, TNLG and technology/sector placebos) and
  logs of unsuccessful pulls (OWID, IMF AIPI, OECD ICT_BUS).

02_industry_accounts/
  Structural business statistics and national accounts for NACE M71, F and M;
  ILOSTAT ISIC F/M employment; BLS engineering comparison (NAICS 54, too broad).

03_trade/
  OECD–WTO BaTIS UK–India and UK–China corridors (SJ3, SE, SI), broader civil-related
  extracts, and the EU importer panels used in the regressions.

04_external_indices/
  Eloundou occupation scores and Study 1 / legacy indices used only as context.
  They are not the generative shock.

05_derived/
  Machine-readable estimates (results*.json) and derived tables matching the article.

06_tables/
  Table1.csv–Table6.csv, the six tables in the main manuscript.

07_scripts/
  Retrieval, estimation and submission-build scripts.

08_figures/
  Figure1.png to Figure5.png (300 dpi), the figures embedded in the manuscript.

Third-party files (Eurostat, OECD–WTO, ILO, ONS, Felten, Eloundou) remain the
property of their publishers and should be cited as in the manuscript reference
list. They are included only as the versions used for the analysis.

Licence for author-generated files (derived tables, figures, scripts, corridor
extracts prepared for this article): CC BY 4.0.
""",
        encoding="utf-8",
    )


def write_readme(path: Path) -> None:
    path.write_text(
        f"""# {JOURNAL} — Study 4 submission set

This folder is the complete journal package for Study 4, in the same role as
`Data_IJCM/` plus the submitted manuscript for Study 1.

## Upload these files to ScholarOne / Taylor & Francis

1. `01_Title_Page_Not_for_Review.docx` — title, author, affiliation, correspondence, declarations. Mark **not for review**.
2. `02_Blinded_Manuscript_for_Review.docx` — anonymous main manuscript (title, abstract, keywords, numbered text, disclosure, data availability, APA 7 references, six tables, five figures). Double-spaced 12-pt Times New Roman, 2.54 cm margins, line and page numbers.
3. `03_Cover_Letter.docx` — letter to the editor.
4. `04_Tables.docx` — editable copies of Tables 1–6.
5. `05_Figure_Captions.docx` — captions for Figures 1–5.
6. `Figures/Figure1.png` … `Figure5.png` — separate 300 dpi files.
7. `Data_Study4_IJCM/` — numbered replication package (Study 1 layout).

An identified author copy is `00_Manuscript_as_Submitted.docx`.

## Journal format applied

- Research article for the *{JOURNAL}* (Taylor & Francis, ISSN 1562-3599 / 2331-2327).
- Unstructured abstract (about 230 words) and six keywords.
- Numbered sections: Introduction; Literature review; Data; Empirical design; Results; Discussion; Limitations; Conclusion.
- End matter required by recent IJCM papers: Disclosure statement; Data availability statement; References.
- In-text citations and reference list in APA 7th author–date form (accepted by Taylor & Francis format-free / Your Paper Your Way; production will apply the journal template after acceptance).
- Tables numbered Table 1–Table 6 with titles above and notes below.
- Figures numbered Figure 1–Figure 5, 300 dpi, captions listed separately.
- Double-anonymous review: the blinded file contains no author name, email, affiliation or acknowledgements.

## Required human check

Confirm funding, competing-interest and originality statements in ScholarOne if any detail has changed.

## Licence

Author-generated files: CC BY 4.0. Third-party statistics remain under their publishers’ terms.
""",
        encoding="utf-8",
    )


def write_inventory(folder: Path) -> None:
    files = sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix != ".zip")
    (folder / "PACKAGE_INVENTORY.md").write_text(
        "# Study 4 IJCM submission inventory\n\n"
        + "\n".join(f"- `{p.relative_to(folder).as_posix()}`" for p in files)
        + "\n",
        encoding="utf-8",
    )
    checksum_files = sorted(
        p for p in folder.rglob("*")
        if p.is_file() and p.name not in {"SHA256SUMS.txt"} and p.suffix != ".zip"
    )
    (folder / "SHA256SUMS.txt").write_text(
        "\n".join(
            f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(folder).as_posix()}"
            for p in checksum_files
        )
        + "\n",
        encoding="utf-8",
    )


def compile_html(md_path: Path, dest: Path) -> None:
    try:
        import markdown
    except ImportError:
        return
    css = """
    body { font-family: "Times New Roman", Times, serif; max-width: 48rem;
           margin: 2rem auto; padding: 0 1.2rem; line-height: 1.6; color: #111; }
    img { max-width: 100%; height: auto; display: block; margin: 0.6rem 0 1.4rem; }
    table { border-collapse: collapse; width: 100%; font-size: 0.9rem; margin: 1rem 0 1.4rem; }
    th, td { border: 1px solid #444; padding: 0.3rem 0.45rem; text-align: left; }
    th { background: #f3f3f3; }
    h1 { font-size: 1.45rem; text-align: center; }
    h2 { font-size: 1.15rem; }
    """
    text = md_path.read_text(encoding="utf-8")
    for dest_name, source_name, _ in MAIN_FIGURES:
        text = text.replace(f"figures/{source_name}", f"Figures/{dest_name}")
    body = markdown.markdown(
        text,
        extensions=["tables", "fenced_code", "nl2br"],
    )
    dest.write_text(
        "<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>"
        f"<title>{TITLE}</title><style>{css}</style></head><body>{body}</body></html>",
        encoding="utf-8",
    )


def word_count(text: str) -> int:
    return len(re.findall(r"[A-Za-z0-9Δ×β−-]+", text))


def main() -> None:
    identified = format_source_manuscript()
    wc = word_count(identified)
    abstract_words = word_count(ABSTRACT)

    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "Figures").mkdir(parents=True)
    (OUT / "figures").mkdir(parents=True)
    for _, source_name, _ in MAIN_FIGURES:
        shutil.copy2(STUDY / "figures" / source_name, OUT / "figures" / source_name)

    build_data_package(DATA_ROOT)
    shutil.copytree(DATA_ROOT, OUT / "Data_Study4_IJCM")
    export_figures(OUT / "Figures")

    reference = OUT / ".ijcm-reference.docx"
    make_reference_doc(reference)

    identified_md = OUT / "00_Manuscript_as_Submitted.md"
    identified_md.write_text(identified, encoding="utf-8")
    pandoc(identified_md, OUT / "00_Manuscript_as_Submitted.docx", reference, OUT)
    polish_docx(OUT / "00_Manuscript_as_Submitted.docx", line_numbers=False, header=RUNNING_HEAD)

    blinded = blinded_markdown(identified)
    blinded_md = OUT / "02_Blinded_Manuscript_for_Review.md"
    blinded_md.write_text(blinded, encoding="utf-8")
    if re.search(r"(?i)yuxuan chai|strathclyde|yuxuanchai98", blinded):
        raise SystemExit("Blinded manuscript still contains identifying information")
    pandoc(blinded_md, OUT / "02_Blinded_Manuscript_for_Review.docx", reference, OUT)
    polish_docx(
        OUT / "02_Blinded_Manuscript_for_Review.docx",
        line_numbers=True,
        header=RUNNING_HEAD,
    )

    write_title_page(OUT / "01_Title_Page_Not_for_Review.docx", wc)
    write_cover_letter(OUT / "03_Cover_Letter.docx")
    write_tables_docx(OUT / "04_Tables.docx")
    write_figure_captions_docx(OUT / "05_Figure_Captions.docx")
    write_readme(OUT / "README_SUBMISSION.md")
    compile_html(blinded_md, OUT / "02_Blinded_Manuscript_for_Review.html")
    compile_html(identified_md, OUT / "00_Manuscript_as_Submitted.html")
    shutil.rmtree(OUT / "figures", ignore_errors=True)
    reference.unlink()
    (OUT / "FORMAT_CHECK.txt").write_text(
        "\n".join(
            [
                f"Journal: {JOURNAL}",
                f"Title: {TITLE}",
                f"Abstract words: {abstract_words}",
                f"Manuscript words: {wc}",
                f"Keywords: {KEYWORDS}",
                "Tables: 6",
                "Figures: 5 at 300 dpi",
                "Blinded identity check: passed",
                "Spacing: double; font: 12-pt Times New Roman; margins: 2.54 cm",
                "",
            ]
        ),
        encoding="utf-8",
    )
    write_inventory(OUT)

    zip_inside = OUT / "Study4_IJCM_Submission.zip"
    zip_root = ROOT / "Study4_IJCM_Submission.zip"
    for path in (zip_inside, zip_root):
        if path.exists():
            path.unlink()
    subprocess.check_call(
        ["zip", "-r", "-X", str(zip_root), ".", "-x", "*.zip"],
        cwd=OUT,
        stdout=subprocess.DEVNULL,
    )
    shutil.copy2(zip_root, zip_inside)
    subprocess.check_call(["unzip", "-t", str(zip_inside)], stdout=subprocess.DEVNULL)
    print(OUT)
    print("abstract_words", abstract_words)
    print("word_count", wc)
    print(zip_inside)


if __name__ == "__main__":
    main()
